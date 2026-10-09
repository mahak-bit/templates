# -*- coding: utf-8 -*-
import os, re
from docx import Document
from docx.shared import Pt, RGBColor, Mm
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from prompts import CATEGORIES, CORE_SIX, FACT_GUARD, expanded
import matter as M
ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
COB=RGBColor(0x00,0x2F,0xA7); INK=RGBColor(0x16,0x16,0x16)
PH=re.compile(r"(\[[A-Z][A-Z0-9 &/'’,\-\.\(\)]*\]|\[(?:NEEDS INFO|CHECK):[^\]]*\])")

def new_doc():
    d=Document(); s=d.sections[0]
    s.page_width,s.page_height=Mm(210),Mm(297)
    for a in ('left_margin','right_margin'): setattr(s,a,Mm(20))
    s.top_margin=Mm(20); s.bottom_margin=Mm(20)
    st=d.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10.5); st.font.color.rgb=INK
    for n,sz in (('Heading 1',26),('Heading 2',16),('Heading 3',12)):
        h=d.styles[n]; h.font.name='Georgia'; h.font.size=Pt(sz); h.font.bold=True; h.font.color.rgb=INK
        h.element.rPr.rFonts.set(qn('w:eastAsia'),'Georgia')
    return d
def shade(cell,hexcol):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd')
    sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),hexcol); tcPr.append(sh)
def runs(p,text,size=None,bold=False,italic=False):
    for part in PH.split(text):
        if not part: continue
        r=p.add_run(part); r.bold=bold; r.italic=italic
        if size: r.font.size=Pt(size)
        if PH.fullmatch(part): r.font.color.rgb=COB; r.bold=True
def para(d,text,**k):
    p=d.add_paragraph(); runs(p,text,**k); return p
def clean(t): return t.replace('**','')
def pb(d): d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
def h(d,t,l=1): return d.add_heading(t,l)
def box(d,text):
    t=d.add_table(rows=1,cols=1); c=t.rows[0].cells[0]; shade(c,'F1F0EE')
    c.paragraphs[0].text=''
    for i,line in enumerate(text.split('\n')):
        p=c.paragraphs[0] if i==0 else c.add_paragraph()
        p.paragraph_format.space_after=Pt(0); runs(p,line,size=10)
    d.add_paragraph()
def brief(d, standalone=False):
    B=M.BRIEF
    h(d,B['heading'],1); para(d,B['lead'],italic=True)
    h(d,B['core_six'],2); para(d,B['core_note'])
    t=d.add_table(rows=0,cols=2); t.style='Table Grid'
    for k,desc in CORE_SIX:
        r=t.add_row().cells; runs(r[0].paragraphs[0],k,bold=True); r[1].paragraphs[0].add_run(desc).font.size=Pt(9)
        r[1].add_paragraph('\n')
    for letter,name,fields in B['sections']:
        h(d,f'{letter}. {name}',2)
        t=d.add_table(rows=0,cols=2); t.style='Table Grid'
        for q in fields:
            r=t.add_row().cells; r[0].text=q; r[0].paragraphs[0].runs[0].font.size=Pt(9.5); r[1].text=''
            r[0].width=Mm(75); r[1].width=Mm(95)
    h(d,B['context_heading'],2); para(d,B['context_lead']); box(d,B['context_template'])

def build_kit():
    d=new_doc()
    h(d,M.TITLE,1); para(d,M.SUBTITLE,size=14,italic=True); para(d,M.AUDIENCE_LINE); para(d,f'{M.VERSION} · {M.YEAR}')
    para(d,'Editable Word version. Fonts used here are Georgia and Calibri so it opens cleanly anywhere; the PDF uses Playfair Display and Inter.',size=9,italic=True)
    pb(d)
    h(d,M.LICENCE['heading'],1)
    for k,v in M.LICENCE['items']:
        p=d.add_paragraph(); r=p.add_run(k+'. '); r.bold=True; p.add_run(v)
    pb(d)
    I=M.INTRO; h(d,I['heading'],1); para(d,I['lead'],italic=True,size=12)
    for t in I['paras']: para(d,clean(t))
    h(d,I['for_heading'],2)
    for a,b in I['for_items']: d.add_paragraph(f'{a} {b}',style='List Bullet')
    h(d,I['inside_heading'],2)
    for t in I['inside_items']: d.add_paragraph(clean(t),style='List Bullet')
    h(d,I['honest_heading'],2); para(d,I['honest'])
    pb(d)
    Q=M.QUICKSTART; h(d,Q['heading'],1); para(d,Q['lead'],italic=True)
    for i,(t,tm,x) in enumerate(Q['steps'],1):
        p=d.add_paragraph(); r=p.add_run(f'{i}. {t} ({tm})'); r.bold=True
        para(d,clean(x))
    h(d,Q['followups_heading'],2)
    for t in Q['followups']: d.add_paragraph(t,style='List Bullet')
    h(d,Q['habits_heading'],2)
    for t in Q['habits']: d.add_paragraph(t,style='List Bullet')
    F=M.FINDER; h(d,F['heading'],1)
    t=d.add_table(rows=0,cols=2); t.style='Table Grid'
    for a,b in F['rows']:
        r=t.add_row().cells; r[0].text=a; r[1].text='Prompts '+b
    h(d,F['order_heading'],2)
    for a,b in F['order']: d.add_paragraph(f'{a}: {b}',style='List Number')
    A=M.ANATOMY; h(d,A['heading'],1); para(d,A['lead'])
    for a,b in A['parts']:
        p=d.add_paragraph(); p.add_run(a+': ').bold=True; p.add_run(b)
    h(d,A['fg_heading'],2); para(d,A['fg_lead']); box(d,FACT_GUARD)
    h(d,A['markers_heading'],2)
    for a,b in A['markers']: para(d,f'{a}  {b}')
    para(d,A['placeholder_note'],size=9)
    pb(d); brief(d); pb(d)
    for c in CATEGORIES:
        h(d,f"{c['n']:02d}  {c['title']}",1); para(d,c['blurb'],italic=True)
        para(d,'Suggested order: '+c['order'],size=9.5)
        for p in c['prompts']:
            pb(d); h(d,f"Prompt {p['n']}: {p['title']}",2)
            para(d,'Use it when: '+p['use'],italic=True)
            para(d,'Fill in',bold=True)
            t=d.add_table(rows=0,cols=2); t.style='Table Grid'
            for k,v in p['fill']:
                r=t.add_row().cells; runs(r[0].paragraphs[0],k,bold=True); r[1].text=v
            para(d,'')
            para(d,'The prompt (copy everything in the box)',bold=True); box(d,expanded(p))
            para(d,'Tip: '+p['tip'],size=9.5)
        pb(d)
    # example
    h(d,'Worked Example: Saltgrain Bakehouse (fictional)',1)
    para(d,'Saltgrain Bakehouse and Eastbrook are fictional; all facts were invented for this example. The sample copy is illustrative, not the output of a specific AI tool.',italic=True)
    h(d,'Key facts',2)
    for t in M.EX_FACTS: d.add_paragraph(t,style='List Bullet')
    h(d,'The Core Six, filled in',2)
    t=d.add_table(rows=0,cols=2); t.style='Table Grid'
    for a,b in M.EX_CORE: r=t.add_row().cells; runs(r[0].paragraphs[0],a,bold=True); r[1].text=b
    h(d,'Prompt 1 filled in',2); box(d,M.EX_HERO_FILLED)
    h(d,'Illustrative hero output',2)
    for i,(a,hl,sb,bt,rs) in enumerate(M.EX_HERO_OUT,1):
        para(d,f'{i}. {a}',bold=True); para(d,hl,size=13,bold=True); para(d,sb); para(d,f'Button: {bt}   |   Under button: {rs}',size=9.5)
    para(d,'Recommendation: '+M.EX_HERO_REC)
    S=M.EX_SERVICE; h(d,'Prompt 11 output: service page',2)
    para(d,S['h1'],size=14,bold=True); para(d,S['intro']); para(d,'Who this is for: '+S['for'])
    para(d,'What is included:',bold=True)
    for t in S['included']: d.add_paragraph(t,style='List Bullet')
    para(d,'How it works:',bold=True)
    for t in S['steps']: d.add_paragraph(t,style='List Number')
    para(d,'What it costs: '+S['cost']); para(d,'Button: '+S['cta'])
    h(d,'Prompt 24 output: FAQ',2)
    for q,a in M.EX_FAQ: para(d,q,bold=True); para(d,a)
    h(d,'Prompt 27 output: titles and descriptions',2)
    for n,t,x in M.EX_META:
        para(d,n,bold=True); para(d,f'Title ({len(t)} chars): {t}'); para(d,f'Description ({len(x)} chars): {x}')
    h(d,'Review with the checklist',2)
    for k,t,x in M.EX_REVIEW: para(d,f"{'PASS' if k=='pass' else 'FLAG'} - {t}: {x}")
    h(d,'What to notice',2)
    for t in M.EX_TAKEAWAYS: d.add_paragraph(t,style='List Bullet')
    pb(d)
    C=M.CHECKLIST; h(d,C['heading'],1); para(d,C['lead'],italic=True)
    for name,items in C['groups']:
        h(d,name,2)
        for t in items: para(d,'☐  '+t)
    h(d,C['redflags_heading'],2); para(d,' · '.join(C['redflags'])); para(d,C['redflags_note'],size=9)
    pb(d); K=M.CLOSING; h(d,K['heading'],1); para(d,K['lead'],italic=True,size=12)
    h(d,K['next_heading'],2)
    for a,b in K['next']: d.add_paragraph(f'{a} {b}',style='List Number')
    para(d,K['closing_line'],size=13,italic=True)
    out=os.path.join(ROOT,'dist','The-AI-Website-Copy-Kit-editable.docx'); d.save(out); print('wrote',out)

def build_brief():
    d=new_doc(); h(d,M.TITLE,1); para(d,'Reusable client & business brief (fillable)',italic=True)
    brief(d); out=os.path.join(ROOT,'dist','Client-and-Business-Brief-fillable.docx'); d.save(out); print('wrote',out)

if __name__=='__main__': build_kit(); build_brief()
